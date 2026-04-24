"""Excel + CSV export service.

Generates RTL-aware Excel workbooks and UTF-8 CSV streams from row data.
Used by `app.api.exports` to ship data to accountants.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Single column descriptor: (key, header_label, optional formatter)
ColumnSpec = tuple[str, str] | tuple[str, str, Callable[[Any], Any]]


_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_BODY_FONT = Font(name="Calibri", size=10)
_DATE_FORMAT = "yyyy-mm-dd"
_MONEY_FORMAT = '#,##0.00'


def _coerce(value: Any) -> Any:
    """Convert value into something openpyxl can write directly."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime) and value.tzinfo is not None:
            return value.replace(tzinfo=None)
        return value
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _resolve(row: dict, spec: ColumnSpec) -> Any:
    key = spec[0]
    raw = row.get(key)
    if len(spec) == 3 and callable(spec[2]):
        try:
            raw = spec[2](raw)
        except Exception:
            pass
    return raw


def to_excel(
    rows: Iterable[dict],
    columns: list[ColumnSpec],
    *,
    title: str = "Sheet1",
    rtl: bool = True,
    money_columns: Iterable[str] | None = None,
    date_columns: Iterable[str] | None = None,
) -> bytes:
    """Build an XLSX file in memory and return raw bytes."""
    wb = Workbook()
    ws = wb.active
    safe_title = (title or "Sheet1")[:31] or "Sheet1"
    ws.title = safe_title
    if rtl:
        ws.sheet_view.rightToLeft = True

    money_set = set(money_columns or [])
    date_set = set(date_columns or [])

    # Header row
    for col_idx, spec in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=spec[1])
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Body rows
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, spec in enumerate(columns, start=1):
            value = _coerce(_resolve(row, spec))
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _BODY_FONT
            key = spec[0]
            if key in money_set and isinstance(value, (int, float)):
                cell.number_format = _MONEY_FORMAT
                cell.alignment = Alignment(horizontal="left" if rtl else "right")
            elif key in date_set and isinstance(value, (datetime, date)):
                cell.number_format = _DATE_FORMAT

    # Column widths (simple heuristic)
    for col_idx, spec in enumerate(columns, start=1):
        header_len = len(str(spec[1]))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(40, header_len + 4))

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def to_csv(rows: Iterable[dict], columns: list[ColumnSpec]) -> bytes:
    """Build a UTF-8 CSV (with BOM for Excel compatibility) and return bytes."""
    text_buf = io.StringIO()
    writer = csv.writer(text_buf, quoting=csv.QUOTE_MINIMAL)
    writer.writerow([spec[1] for spec in columns])
    for row in rows:
        writer.writerow([_coerce(_resolve(row, spec)) for spec in columns])
    # Prepend UTF-8 BOM so Excel renders Kurdish/Arabic correctly
    return ("\ufeff" + text_buf.getvalue()).encode("utf-8")


def filename_for(base: str, fmt: str) -> str:
    safe = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in base)
    ext = "xlsx" if fmt == "excel" else "csv"
    return f"{safe}.{ext}"


def media_type_for(fmt: str) -> str:
    if fmt == "excel":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "text/csv; charset=utf-8"


def render(
    rows: Iterable[dict],
    columns: list[ColumnSpec],
    *,
    fmt: str,
    title: str,
    money_columns: Iterable[str] | None = None,
    date_columns: Iterable[str] | None = None,
) -> tuple[bytes, str, str]:
    """Render rows as either Excel or CSV. Returns (bytes, media_type, filename)."""
    fmt = (fmt or "excel").lower()
    if fmt not in {"excel", "csv"}:
        fmt = "excel"
    if fmt == "excel":
        payload = to_excel(
            rows, columns, title=title,
            money_columns=money_columns, date_columns=date_columns,
        )
    else:
        payload = to_csv(rows, columns)
    return payload, media_type_for(fmt), filename_for(title, fmt)
